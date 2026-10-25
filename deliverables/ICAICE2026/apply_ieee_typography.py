# -*- coding: utf-8 -*-
"""按 IEEE 会议模板（Conference-template-letter.docx）字号规范统一中文稿排版。

规范来源（模板样式实测值）：
  标题 24pt / 作者 11pt / 摘要 9pt 加粗 / 关键词 10pt /
  正文 10pt、行距固定 11.4pt、段后 6pt、两端对齐 /
  一级标题 10pt 加粗、二级标题 10pt 斜体 / 图题表题 8pt /
  表格 8pt（表头加粗）/ 参考文献 8pt、行距固定 9pt / 公式 10pt
中西文字体：西文 Times New Roman，中文 宋体（全篇统一，不用黑体/雅黑）。

只改字体、字号、行距、对齐，不改文字内容与图表位置。
"""
import re
import shutil
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"E:\Intrel-Train Industrial\deliverables\ICAICE2026\规则增强的变电运维个性化培训智能体设计与实现_修订版1021_v2.docx"
BAK = r"E:\Intrel-Train Industrial\deliverables\ICAICE2026\_备份_格式调整前_20261021.docx"

WEST = "Times New Roman"
EAST = "宋体"

shutil.copy2(SRC, BAK)
d = docx.Document(SRC)


# ---------- 基础工具 ----------
def set_run(run, size=None, bold=None, italic=None, small_caps=None):
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(attr)) is not None:
            del rf.attrib[qn(attr)]
    rf.set(qn("w:ascii"), WEST)
    rf.set(qn("w:hAnsi"), WEST)
    rf.set(qn("w:eastAsia"), EAST)
    rf.set(qn("w:cs"), WEST)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if small_caps is not None:
        run.font.small_caps = small_caps


def fmt_par(p, align=None, line_exact=None, before=None, after=None,
            indent_chars=None, keep_next=None):
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    if line_exact is not None:
        pf.line_spacing = Pt(line_exact)      # Length -> lineRule=exact
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if indent_chars is not None:
        pPr = p._p.get_or_add_pPr()
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        for a in ("w:firstLine", "w:firstLineChars"):
            if ind.get(qn(a)) is not None:
                del ind.attrib[qn(a)]
        ind.set(qn("w:firstLineChars"), str(int(indent_chars * 100)))
    if keep_next is not None:
        pf.keep_with_next = keep_next


def fmt_all_runs(p, **kw):
    for r in p.runs:
        set_run(r, **kw)


def fmt_math(p, size):
    """公式（OMML）内部 run 字号"""
    for rpr in p._p.iter(qn("m:r")):
        wpr = rpr.find(qn("w:rPr"))
        if wpr is None:
            wpr = OxmlElement("w:rPr")
            rpr.insert(0, wpr)
        for tag in ("w:sz", "w:szCs"):
            e = wpr.find(qn(tag))
            if e is None:
                e = OxmlElement(tag)
                wpr.append(e)
            e.set(qn("w:val"), str(int(size * 2)))


# ---------- 段落分类 ----------
paras = d.paragraphs
texts = [p.text.strip() for p in paras]

abs_i = next(i for i, t in enumerate(texts) if t.startswith("摘要"))
kw_i = next(i for i, t in enumerate(texts) if t.startswith("关键词"))
ref_i = max(i for i, t in enumerate(texts) if t == "参考文献")

H1 = re.compile(r"^\d+\s")
H2 = re.compile(r"^\d+\.\d+(\.\d+)?\s")
CAP = re.compile(r"^[图表]\s*\d")

stats = {"title": 0, "author": 0, "abstract": 0, "keywords": 0,
         "h1": 0, "h2": 0, "caption": 0, "body": 0, "ref": 0, "blank": 0}

for i, p in enumerate(paras):
    t = texts[i]

    if i == 0:                                        # 论文标题
        fmt_all_runs(p, size=24, bold=True)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, line_exact=None)
        stats["title"] += 1
    elif i < abs_i:                                   # 作者 / 单位 / 邮箱
        fmt_all_runs(p, size=11)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=2)
        stats["author"] += 1
    elif i == abs_i:                                  # 摘要
        fmt_all_runs(p, size=9, bold=True)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=6, after=10)
        stats["abstract"] += 1
    elif i == kw_i:                                   # 关键词
        fmt_all_runs(p, size=10)
        if p.runs:                                    # 标签加粗，词条正常
            set_run(p.runs[0], size=10, bold=True)
            for r in p.runs[1:]:
                set_run(r, size=10, bold=False)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6)
        stats["keywords"] += 1
    elif t == "参考文献":                              # 参考文献标题
        fmt_all_runs(p, size=10, bold=True)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=4, keep_next=True)
        stats["h1"] += 1
    elif i > ref_i and t:                             # 参考文献条目
        fmt_all_runs(p, size=8)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=9, before=0, after=2.5)
        stats["ref"] += 1
    elif i > ref_i:
        stats["blank"] += 1
    elif H1.match(t) and len(t) < 30:                 # 一级标题
        fmt_all_runs(p, size=10, bold=True, italic=False, small_caps=True)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_exact=None,
                before=8, after=4, indent_chars=0, keep_next=True)
        stats["h1"] += 1
    elif H2.match(t) and len(t) < 30:                 # 二级标题
        fmt_all_runs(p, size=10, bold=False, italic=True)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_exact=None,
                before=6, after=3, indent_chars=0, keep_next=True)
        stats["h2"] += 1
    elif CAP.match(t) and len(t) < 40:                # 图题 / 表题
        fmt_all_runs(p, size=8)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=6)
        stats["caption"] += 1
    elif t:                                           # 正文
        fmt_all_runs(p, size=10)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=11.4,
                before=0, after=6, indent_chars=2)
        stats["body"] += 1
    else:
        stats["blank"] += 1

# ---------- 公式段落 ----------
for p in paras:
    if p._p.findall(".//" + qn("m:oMath")) or p._p.findall(".//" + qn("m:oMathPara")):
        fmt_all_runs(p, size=10)
        fmt_math(p, 10)
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6, indent_chars=0)

# ---------- 表格 ----------
for t in d.tables:
    for ri, row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run(r, size=8, bold=True if ri == 0 else None)
                p.paragraph_format.line_spacing = Pt(9)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)

# ---------- 文档默认样式（Normal）----------
normal = d.styles["Normal"]
normal.font.size = Pt(10)
normal.font.name = WEST
rpr = normal.element.get_or_add_rPr()
rf = rpr.find(qn("w:rFonts"))
if rf is None:
    rf = OxmlElement("w:rFonts")
    rpr.insert(0, rf)
rf.set(qn("w:ascii"), WEST)
rf.set(qn("w:hAnsi"), WEST)
rf.set(qn("w:eastAsia"), EAST)

d.save(SRC)
print("saved:", SRC)
print("统计:", stats)

# ---------- 自检 ----------
d2 = docx.Document(SRC)
from collections import Counter
c = Counter()
for p in d2.paragraphs:
    for r in p.runs:
        c[(r.font.size.pt if r.font.size else None,
           r.font.name,
           r._r.rPr.rFonts.get(qn("w:eastAsia")) if r._r.rPr is not None and r._r.rPr.rFonts is not None else None)] += 1
print("--- 正文 run 字体/字号分布 ---")
for k, v in c.most_common(10):
    print(f"  {v:4d}  {k}")
print("--- 抽查 ---")
for i in (0, 1, 4, 5, 7, 8, 9, 17, 24, 81, 82):
    p = d2.paragraphs[i]
    pf = p.paragraph_format
    print(f"  [{i}] {p.text[:26]!r} size={p.runs[0].font.size.pt if p.runs else None} "
          f"bold={p.runs[0].font.bold if p.runs else None} "
          f"italic={p.runs[0].font.italic if p.runs else None} "
          f"align={p.alignment} line={pf.line_spacing} after={pf.space_after}")
