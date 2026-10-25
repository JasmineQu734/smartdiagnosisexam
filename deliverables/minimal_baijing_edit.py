from pathlib import Path
import shutil

from docx import Document


SRC = Path(r"D:\intelligent-learning-system-develop\intelligent-learning-system-develop\百景大赛开放创新赛道-慧诊智考项目报告.docx")
OUT = Path(r"E:\Intrel-Train Industrial\deliverables\百景大赛开放创新赛道-慧诊智考项目报告-最小修订版.docx")

TRACK_NOTE = (
    "官网公开配置将开放创新赛道定义为“不设固定赛题，鼓励自主挖掘工业AI需求场景”，"
    "本作品即以变电运维人才培训与考核为自主挖掘的工业应用场景。"
)

def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))

    # 1.2 段落内补比赛官网对开放创新赛道的定义和本作品对应关系。
    for para in doc.paragraphs:
        if para.text.strip().startswith("本作品对应开放创新赛道中"):
            if TRACK_NOTE not in para.text:
                para.text = para.text + TRACK_NOTE
            break
    else:
        raise RuntimeError("未找到 1.2 赛道定位段落")

    # 3.2 节只修正实现边界：个人/公司诊断严格契约校验；报告归因支持 model 并失败回退。
    for para in doc.paragraphs:
        if para.text.strip().startswith("三类智能服务（报告归因、个人个性化诊断、公司运行诊断）走同一套可信诊断闭环"):
            para.text = para.text.replace(
                "三类智能服务（报告归因、个人个性化诊断、公司运行诊断）走同一套可信诊断闭环",
                "个人个性化诊断与公司运行诊断走同一套可信诊断闭环",
            )
            para.text = para.text.replace(
                "接口兼容OpenAI与Anthropic两种协议",
                "个人与公司诊断接口兼容OpenAI与Anthropic两种协议",
            )
            para.text = para.text.replace(
                "任何不合格都视为本次生成失败。",
                "任何不合格都视为本次生成失败。报告归因另支持source=rule或source=model选路，模型失败回退规则并写入fallback_note。",
            )
            break
    else:
        raise RuntimeError("未找到 3.2 大模型协同段落")

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
