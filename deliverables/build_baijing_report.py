from pathlib import Path
import textwrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
    KeepTogether,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
DOCX = OUT / "百景大赛开放创新赛道-慧诊智考项目报告-修订版.docx"
PDF = OUT / "百景大赛开放创新赛道-慧诊智考项目报告-修订版.pdf"


TITLE = "慧诊智考项目报告"
SUBTITLE = "面向百景大赛开放创新赛道的可信工业 AI 应用方案"
TEAM = "逐光造梦"
DATE = "2026 年 10 月 6 日"


def font_path():
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/arphic/uming.ttc"),
    ]
    for p in candidates:
        if p.exists():
            return p
    return None


FONT = font_path()
FONT_NAME = "CJK"
if FONT:
    pdfmetrics.registerFont(TTFont(FONT_NAME, str(FONT), subfontIndex=0))
else:
    FONT_NAME = "Helvetica"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(9.5)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def doc_add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        shade_cell(hdr[i], "1F4E79")
        set_cell_text(hdr[i], h, True)
        for run in hdr[i].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    doc.add_paragraph()
    return table


def p(text):
    return textwrap.dedent(text).strip()


sections = [
    {
        "title": "作品简介",
        "paras": [
            p(
                """
                慧诊智考是一套面向变电运维中心的可信工业 AI 应用系统，聚焦“培训有记录、考核有证据、
                诊断可追溯、建议需确认”的一线管理问题。系统覆盖员工、班组长、企业管理者和平台管理员
                四类角色，把技能画像、个性化诊断、在线考试、实操报告归因、资料审核和公司级人才驾驶舱
                放在同一套业务闭环里。
                """
            ),
            p(
                """
                本次修订按百景大赛官网公开配置重新定位：2026 百景大赛的赛题公示包含工艺场景赛道、
                通用场景赛道和开放创新赛道，其中开放创新赛道“不设固定赛题，鼓励自主挖掘工业 AI 需求场景”。
                本项目并非套用固定赛题，而是从强安全行业的培训考核、能力诊断和报告归因中自主挖掘工业 AI
                需求，符合开放创新赛道的参赛逻辑。
                """
            ),
            p(
                """
                系统当前以模拟数据集完成全流程验证：86 名员工、9 个班组、53 张 MySQL 表、11 个幂等初始化脚本。
                结论来源坚持“规则守下限、模型提上限”：未配置模型或模型输出不合约束时，业务仍走规则引擎；
                配置模型后，个人诊断、公司诊断和报告归因按事实摘要生成结论，并通过契约校验与人工确认闭环。
                """
            ),
        ],
        "tables": [
            (
                ["核对项", "原 PDF 状态", "修订处理"],
                [
                    ["赛道定位", "已写为开放创新赛道，但赛道要求引用不够直接", "补充官网公开配置中的开放创新赛道定义"],
                    ["目录", "1.2 目录项交叉引用失效", "重排结构，不再使用失效交叉引用"],
                    ["Agent 状态", "容易把所有 Agent 统称为已完整接模型", "区分业务诊断入口与通用 Agent 契约桩"],
                    ["证据边界", "已说明模拟数据，但分散在后文", "在简介、测试和前景中重复明确“未企业试点”"],
                ],
            )
        ],
    },
    {
        "title": "一 项目与赛道适配",
        "paras": [
            p(
                """
                开放创新赛道看重参赛者能否自主发现工业现场的真实 AI 需求，并把方案做成可验证、可迁移、
                可落地的工程系统。本项目的需求来自变电运维中心的真实管理痛点：培训台账分散、岗位技能底数
                不清、考试与实操割裂、报告归因依赖人工经验、关键岗位备份不足。
                """
            ),
            p(
                """
                与普通学习管理系统相比，本项目的工业属性体现在三处。第一，判断依据来自岗位技能等级、
                国家标准阈值、考试记录、实操报告和证书有效期，而不是泛化学习进度。第二，安全不合格、
                有限空间气体超限、未按 SOP 执行等事项会影响人才决策，不能由模型自由发挥。第三，结论必须
                能回查到规则编号、事实编号、人工确认记录和审计日志。
                """
            ),
        ],
        "tables": [
            (
                ["开放创新赛道要求", "本项目对应设计"],
                [
                    ["自主挖掘工业 AI 场景", "从变电运维培训、考试、实操报告和人才驾驶舱中抽象需求"],
                    ["不设固定赛题", "不依赖固定赛题数据，使用可迁移的岗位技能与规则引擎架构"],
                    ["强调应用价值", "面向降本、补短板、合规审计和关键岗位备份提供闭环"],
                    ["强调可复制", "单一 MySQL、Spring Boot、Vue 和离线规则兜底，便于基层部署"],
                ],
            )
        ],
    },
    {
        "title": "二 系统总体设计",
        "paras": [
            p(
                """
                系统采用“前端单页应用、后端服务、规则与模型协同、单一数据库”的四层结构。前端为 Vue 3、
                Vite 与 Element Plus，使用相对路径 /api/v1 调用后端；后端为 Spring Boot 3，数据访问采用
                JdbcTemplate 和统一 Sql 助手，不引入 ORM；数据层只使用 MySQL 8，一套库同时保存主数据、
                业务数据、规则、Agent 结果和审计留痕。
                """
            ),
            p(
                """
                四类角色的数据边界保持清晰：员工只看本人，班组长看本人和本班组，企业管理者看全公司汇总，
                管理员负责平台配置和审计。心理与关怀数据默认本人可见，管理端只返回脱敏聚合；安全不合格
                在人才决策中优先级最高。
                """
            ),
        ],
        "tables": [
            (
                ["角色", "演示账号", "数据范围", "核心能力"],
                [
                    ["员工", "zhangwei", "本人", "技能雷达、记忆、学习、考试、个人诊断"],
                    ["班组长", "limin", "本人和本班组", "员工功能、报告填报、报告台账、资料导入"],
                    ["企业管理者", "wangzong", "全公司汇总", "公司总览、技能矩阵、培训看板、运行诊断"],
                    ["管理员", "admin", "平台数据", "用户权限、课程题库、资料审核、考试配置、审计"],
                ],
            ),
            (
                ["数据对象", "实现依据"],
                [
                    ["主数据", "employee、org、position、skill、position_skill_requirement"],
                    ["考试与练习", "exam、exam_history、question、practice_session、practice_stat"],
                    ["报告与归因", "task、task_report、rule、rule_result、agent_suggestion、decision_log"],
                    ["诊断与记忆", "diagnosis_result、memory_item、dataset_meta、dataset_item"],
                ],
            ),
        ],
    },
    {
        "title": "三 核心算法与可信机制",
        "paras": [
            p(
                """
                规则引擎是系统可信下限。规则存放在 rule 表，条件只允许白名单算子，不使用 SpEL、eval、
                脚本引擎或拼 SQL。REPORT_GAS 场景已把原先写在 Java 中的气体归因逻辑迁入数据库规则：
                12 条规则分为定档规则和关键词补充规则，支持 rule_group、stop_on_match 和 max_hits，
                同一输入可复算出稳定的 rule_refs。
                """
            ),
            p(
                """
                大模型是增强层，不是业务数据的直接操作者。个人诊断和公司诊断先从 MySQL 读取事实清单，
                再把事实摘要交给模型；模型返回 JSON 后必须通过字段、枚举和数字出处校验，任何不合格都会
                回退到内置规则结论，并写入 fallback_note。报告归因同样支持 source=rule 或 source=model，
                模型不可用时自动回退规则路径。
                """
            ),
            p(
                """
                需要特别说明的是，系统仍保留 POST /api/v1/agent/personal/recommend、
                POST /api/v1/agent/management/decision 和 POST /api/v1/agent/report/analyze 这类通用 Agent
                契约桩；它们当前返回 source=rule 的固定建议并写入 agent_suggestion，主要用于保持未来模型
                替换时的接口形状稳定。本报告的“实时诊断”运行效果来自 /personal/agent-diagnose、
                /exec/agent-diagnose 和 /reports/{reportId}/analyze 等业务入口，不混同为通用桩已完整接模型。
                """
            ),
        ],
        "tables": [
            (
                ["机制", "解决的问题", "当前实现"],
                [
                    ["声明式规则", "业务判断不可黑盒", "rule / rule_result，命中规则和事实依据留痕"],
                    ["契约校验", "模型 JSON 不稳定", "DiagnosisContract 逐字段校验，失败即回退"],
                    ["双源回退", "AI 不可用时业务中断", "AI_API_KEY 为空或失败时 source=rule"],
                    ["人工确认", "建议与执行混同", "agent_suggestion 待确认，decision_log 记录采纳/驳回/修改"],
                ],
            )
        ],
    },
    {
        "title": "四 运行效果与实现证据",
        "paras": [
            p(
                """
                当前仓库已经实现从主数据、考试、个人系统、管理板块、实操任务到规则引擎的主链路。报告填报
                使用统一框格模型，后端字段校验失败返回 code=40001 和 field_errors，校验通过后由 PDFBox 生成
                PDF、Apache POI XWPF 生成 Word。报告归因只读取报告本体正文，不把上一轮归因结论回灌为新事实。
                """
            ),
            p(
                """
                Windows 免安装包采用 jlink 精简 JRE 与绿色版 MySQL，前端由后端托管为同源单页应用。发布包中的
                .env 模板保留 AI_API_KEY= 空值，收件人自行填写密钥；未填写时系统按规则兜底运行，不会为了
                演示而夹带或虚构模型凭据。
                """
            ),
        ],
        "tables": [
            (
                ["指标", "当前口径"],
                [
                    ["初始化脚本", "001 至 011 共 11 个脚本，支持幂等补跑"],
                    ["数据库表", "db/init 定义 53 张表"],
                    ["演示组织", "86 名员工、9 个班组、5 类核心技能"],
                    ["实操报告", "演示数据含报告台账，并支持实时 task_report"],
                    ["规则场景", "PERSONAL_GAP、EXAM_RESULT、TALENT_GAP、REPORT_GAS 等"],
                ],
            ),
            (
                ["核对结论", "证据"],
                [
                    ["报告 PDF 的项目方向与当前仓库基本一致", "Spring Boot、Vue、MySQL、规则引擎、报告归因等均有源码与脚本支撑"],
                    ["需要修正旧 PDF 中目录书签错误", "原目录 1.2 的交叉引用已经失效"],
                    ["需要明确未企业试点", "当前数据为演示/模拟口径，不能写成企业上线成效"],
                    ["需要明确通用 Agent 桩边界", "ApiController 中 /agent/{agentType}/{action} 仍返回固定建议"],
                ],
            ),
        ],
    },
    {
        "title": "五 测试验证与边界声明",
        "paras": [
            p(
                """
                本项目的验证重点不是追求大模型回答漂亮，而是证明“事实能读到、规则能复算、模型失败能回退、
                建议不会越权改业务数据”。因此测试分为建库、构建、规则、鉴权和端到端留痕五层。规则引擎专项
                需覆盖启用停用、优先级、分组、组内命中上限、*_from 引用事实、无匹配兜底、脏数据容错和
                rule_code 输出。
                """
            ),
            p(
                """
                边界必须写清：当前演示数据、成本节省和训练规模属于模拟数据集或算法测算，不等同于企业试点成效。
                后续参加百景大赛时，可把“试点单位脱敏数据接入、内网部署、真实用户反馈”作为路演后的下一阶段
                目标，而不是在书面报告中提前承诺已经完成。
                """
            ),
        ],
        "tables": [
            (
                ["验证层", "通过判据"],
                [
                    ["建库", "11 个脚本空库按序执行，表数与定义一致"],
                    ["构建", "后端 Maven 打包、前端构建或既有包可运行"],
                    ["规则", "REPORT_GAS 同一输入 rule_refs 稳定，停用规则不参与"],
                    ["鉴权", "非管理员访问 /rules/** 返回 403 / code=40301"],
                    ["留痕", "rule_result、agent_suggestion、decision_log 可用同一 run_id 串联"],
                ],
            )
        ],
    },
    {
        "title": "六 创新点与推广建议",
        "paras": [
            p(
                """
                本项目的第一项创新是把大模型能力放入可控业务闭环，而不是把模型当作最终裁判。规则引擎保证下限，
                输出契约保证结构，人工确认保证治理责任，fallback_note 保证失败可解释。这个组合比“直接调用大模型
                写诊断报告”更适合强安全行业。
                """
            ),
            p(
                """
                第二项创新是规则即数据。判定逻辑不藏在代码里，而是进入数据库规则表，以 scene、rule_group、
                priority、max_hits 和 action_json 表达。业务口径调整时可以改规则与模板，不需要重写 Controller
                或前端接口。
                """
            ),
            p(
                """
                第三项创新是可迁移的岗位能力治理范式。变电运维只是落地样本，同一套“岗位技能矩阵、实操报告、
                证据链诊断、人工确认、审计留痕”可以迁移到新能源场站、智能制造产线、安全生产培训和设备运维质控。
                对百景大赛而言，项目后续最需要的是开放真实工业场景、脱敏数据窗口和内网试点环境，以便把当前
                模拟验证推进到企业闭环验证。
                """
            ),
        ],
        "tables": [
            (
                ["传统方式痛点", "慧诊智考方案", "可推广价值"],
                [
                    ["培训只留进度", "能力、证书、考试、实操一体化入库", "培训投入可按缺口投放"],
                    ["AI 结论不可控", "模型输出需过契约校验，失败回退规则", "适合强监管行业"],
                    ["报告归因靠经验", "规则编号、事实依据、人工确认串联", "降低复核和责任界定成本"],
                    ["部署门槛高", "单库、双进程、离线可用", "基层班组可试点"],
                ],
            )
        ],
    },
]


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)

    styles = doc.styles
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(10.5)
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.bold = True

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(TITLE)
    run.font.size = Pt(22)
    run.bold = True
    sub = doc.add_paragraph(SUBTITLE)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph(f"团队名称：{TEAM}    赛道名称：开放创新赛道    日期：{DATE}")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("作品名称：慧诊智考——基于可信 AI 的变电运维人才培养诊断与考试管理平台").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    doc.add_paragraph("修订说明：本版根据当前项目仓库、用户提供 PDF 与百景大赛官网公开配置重新核对生成，原 PDF 不覆盖。")
    doc.add_page_break()

    toc = doc.add_paragraph()
    toc.add_run("目录").bold = True
    for i, s in enumerate(sections, 1):
        doc.add_paragraph(f"{i}. {s['title']}")
    doc.add_page_break()

    for s in sections:
        doc.add_heading(s["title"], level=1)
        for para in s["paras"]:
            doc.add_paragraph(para)
        for headers, rows in s.get("tables", []):
            doc_add_table(doc, headers, rows)

    doc.add_heading("附录 资料来源与提交提示", level=1)
    for para in [
        "百景大赛官网：https://100aicv.com/ 。本次核对读取公开配置接口 /api/config/published/site_competition/competition_tracks_list，开放创新赛道描述为“不设固定赛题，鼓励自主挖掘工业AI需求场景”。",
        "项目仓库核对范围：README.md、docs/api-catalog.md、db/init、backend/src/main/java/com/power/train、frontend/src。",
        "提交前建议补充报名后生成的参赛编号，并在路演材料中另附系统截图或演示视频；本报告不把模拟数据写成企业试点结果。",
    ]:
        doc.add_paragraph(para)
    doc.save(DOCX)


def styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("title", parent=base["Title"], fontName=FONT_NAME, fontSize=24, leading=32, alignment=TA_CENTER, textColor=colors.black),
        "subtitle": ParagraphStyle("subtitle", parent=base["Normal"], fontName=FONT_NAME, fontSize=13, leading=20, alignment=TA_CENTER),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontName=FONT_NAME, fontSize=15, leading=22, spaceBefore=14, spaceAfter=8, textColor=colors.black),
        "body": ParagraphStyle("body", parent=base["BodyText"], fontName=FONT_NAME, fontSize=9.5, leading=16, firstLineIndent=18, spaceAfter=6),
        "small": ParagraphStyle("small", parent=base["BodyText"], fontName=FONT_NAME, fontSize=8.5, leading=13, spaceAfter=4),
    }


def pdf_table(headers, rows):
    data = [[Paragraph(h, PDF_STYLES["small"]) for h in headers]]
    for row in rows:
        data.append([Paragraph(str(c), PDF_STYLES["small"]) for c in row])
    col_width = (A4[0] - 4.2 * cm) / len(headers)
    table = Table(data, colWidths=[col_width] * len(headers), repeatRows=1, hAlign="CENTER")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9D9D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def add_page_number(canvas, doc):
    canvas.saveState()
    canvas.setFont(FONT_NAME, 8)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawCentredString(A4[0] / 2, 1.1 * cm, str(doc.page))
    canvas.restoreState()


def build_pdf():
    story = []
    story.append(Spacer(1, 4.5 * cm))
    story.append(Paragraph(TITLE, PDF_STYLES["title"]))
    story.append(Spacer(1, 0.5 * cm))
    story.append(Paragraph(SUBTITLE, PDF_STYLES["subtitle"]))
    story.append(Spacer(1, 1.0 * cm))
    story.append(Paragraph("作品名称：慧诊智考——基于可信 AI 的变电运维人才培养诊断与考试管理平台", PDF_STYLES["subtitle"]))
    story.append(Paragraph(f"团队名称：{TEAM}    赛道名称：开放创新赛道", PDF_STYLES["subtitle"]))
    story.append(Paragraph(f"日期：{DATE}", PDF_STYLES["subtitle"]))
    story.append(PageBreak())
    story.append(Paragraph("目录", PDF_STYLES["h1"]))
    for i, s in enumerate(sections, 1):
        story.append(Paragraph(f"{i}. {s['title']}", PDF_STYLES["body"]))
    story.append(PageBreak())

    for s in sections:
        story.append(Paragraph(s["title"], PDF_STYLES["h1"]))
        for para in s["paras"]:
            story.append(Paragraph(para, PDF_STYLES["body"]))
        for headers, rows in s.get("tables", []):
            story.append(KeepTogether([pdf_table(headers, rows), Spacer(1, 0.25 * cm)]))

    story.append(Paragraph("附录 资料来源与提交提示", PDF_STYLES["h1"]))
    for para in [
        "百景大赛官网：https://100aicv.com/ 。本次核对读取公开配置接口 /api/config/published/site_competition/competition_tracks_list，开放创新赛道描述为“不设固定赛题，鼓励自主挖掘工业AI需求场景”。",
        "项目仓库核对范围：README.md、docs/api-catalog.md、db/init、backend/src/main/java/com/power/train、frontend/src。",
        "提交前建议补充报名后生成的参赛编号，并在路演材料中另附系统截图或演示视频；本报告不把模拟数据写成企业试点结果。",
    ]:
        story.append(Paragraph(para, PDF_STYLES["body"]))

    doc = BaseDocTemplate(
        str(PDF),
        pagesize=A4,
        leftMargin=2.1 * cm,
        rightMargin=2.1 * cm,
        topMargin=1.9 * cm,
        bottomMargin=1.8 * cm,
        title="百景大赛开放创新赛道-慧诊智考项目报告-修订版",
        author=TEAM,
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=add_page_number)])
    doc.build(story)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    PDF_STYLES = styles()
    build_docx()
    build_pdf()
    print(DOCX)
    print(PDF)
