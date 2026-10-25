# -*- coding: utf-8 -*-
"""生成 IEEE 双栏 Word（.docx）。公式以真正的 OMML 对象写入，可在 Word 里直接编辑。"""
from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "main_IEEE.docx")

BY_NUM = {int(t["num"]): t for t in C.TABLES}

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ---------------------------------------------------------------- OMML

def _m(tag, *children, **attrs):
    e = OxmlElement("m:" + tag)
    for k, v in attrs.items():
        e.set(qn("m:" + k), v)
    for c in children:
        if c is not None:
            e.append(c)
    return e


def _t(text, style="p"):
    """math text run。style: p=正体, i=斜体"""
    r = _m("r")
    rPr = OxmlElement("m:rPr")
    if style == "p":
        st = OxmlElement("m:sty"); st.set(qn("m:val"), "p"); rPr.append(st)
    r.append(rPr)
    wt = OxmlElement("w:t")
    wt.text = text
    wt.set(qn("xml:space"), "preserve")
    r.append(wt)
    return r


def omml_from_tokens(tokens):
    """token DSL -> m:oMathPara"""
    def build(t):
        if isinstance(t, str):
            return _t(t, "p")
        kind = t[0]
        if kind == "i":
            return _t(t[1], "i")
        if kind in ("sub", "sup"):
            base_text = ""
            e = _m("sSub" if kind == "sub" else "sSup")
            e.append(_m("e", _t("", "i")))
            sub = _m("sub" if kind == "sub" else "sup")
            for x in t[1]:
                sub.append(build(x))
            e.append(sub)
            return e
        if kind == "frac":
            f = _m("f")
            num = _m("num")
            for x in t[1]:
                num.append(build(x))
            den = _m("den")
            for x in t[2]:
                den.append(build(x))
            f.append(num); f.append(den)
            return f
        raise ValueError(kind)

    omath = _m("oMath")
    for t in tokens:
        omath.append(build(t))
    return _m("oMathPara", omath)


def add_equation(doc, tokens, number: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    # 公式本体
    p._p.append(omml_from_tokens(tokens))
    # 右侧编号（用制表符 + (n)）
    run = p.add_run("\t(%d)" % number)
    run.font.size = Pt(10)
    return p


# ---------------------------------------------------------------- 富文本

BOLD_RE = re.compile(r"\*\*(.+?)\*\*", re.S)
ITAL_RE = re.compile(r"(?<!\*)\*([^*]+?)\*(?!\*)", re.S)


def add_rich(p, text: str, size=10, font="Times New Roman"):
    """支持 **bold** 与 *italic*。"""
    pos = 0
    tokens = []
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            tokens.append(("plain", text[pos:m.start()]))
        tokens.append(("bold", m.group(1)))
        pos = m.end()
    if pos < len(text):
        tokens.append(("plain", text[pos:]))

    def emit(seg, style):
        p2 = 0
        for m in ITAL_RE.finditer(seg):
            if m.start() > p2:
                _run(p, seg[p2:m.start()], style, size=size, font=font)
            _run(p, m.group(1), style, italic=True, size=size, font=font)
            p2 = m.end()
        if p2 < len(seg):
            _run(p, seg[p2:], style, size=size, font=font)

    for style, seg in tokens:
        emit(seg, style)


def _run(p, text, style="plain", italic=False, size=10, font="Times New Roman"):
    if not text:
        return
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.italic = italic
    if style == "bold":
        r.bold = True
    return r


# ---------------------------------------------------------------- 版式工具

def set_two_columns(section, space_twips=360):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), str(space_twips))
    cols.set(qn("w:equalWidth"), "1")


def set_one_column(section):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), "1")


def style_base(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = st.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.0
    pf.first_line_indent = Inches(0.16)



TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
               "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]


def fix_table_layout(tb, width_in):
    """固定布局 + 收紧单元格内边距；并按规定顺序重排 tblPr（顺序错 Word 会忽略）。"""
    tblPr = tb._tbl.tblPr
    for tag in ("tblW", "tblLayout", "tblCellMar"):
        for el in tblPr.findall(qn("w:" + tag)):
            tblPr.remove(el)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(int(width_in * 1440))); w.set(qn("w:type"), "dxa")
    tblPr.append(w)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for side, tw in (("left", 45), ("right", 45), ("top", 10), ("bottom", 10)):
        node = OxmlElement("w:" + side); node.set(qn("w:w"), str(tw)); node.set(qn("w:type"), "dxa")
        mar.append(node)
    tblPr.append(mar)
    rank = {n: i for i, n in enumerate(TBLPR_ORDER)}
    kids = list(tblPr)
    kids.sort(key=lambda el: rank.get(el.tag.split("}")[-1], 99))
    for el in kids:
        tblPr.append(el)


def add_table(doc, t, col_width=3.25):
    n = len(t["header"])
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(2)
    _run(cap, "TABLE %s  " % t["num"], "bold", size=8)
    add_rich(cap, t["caption"], size=8)

    tb = doc.add_table(rows=1, cols=n)
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tb.rows[0].cells
    trPr = tb.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for i, h in enumerate(t["header"]):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, h, "bold", size=7.5)
    for row in t["rows"]:
        cells = tb.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            add_rich(p, str(v), size=7.5)
    for row in tb.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
    # 尽量让整张表与表题落在同一栏：除末行外都设 keep-with-next
    cap.paragraph_format.keep_with_next = True
    for row in tb.rows[:-1]:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.keep_with_next = True

    # 固定列宽，按内容长度（开方压缩）分配，避免规则编号列被挤成竖排
    import math
    from docx.shared import Inches as _In
    weights = []
    for i in range(n):
        longest = max([len(str(t["header"][i]))] + [len(str(r[i])) for r in t["rows"]])
        weights.append(math.sqrt(max(longest, 4)))
    total = sum(weights)
    widths = [max(col_width * w / total, col_width * 0.12) for w in weights]

    # 每列至少装得下最长的不可断西文片段，否则 Word 会在词内折行（如 Knowle/dge）
    def token_in(text, size=7.5):
        best = 0
        for tok in re.split(r"[^A-Za-z0-9_\-.]+", str(text)):
            best = max(best, len(tok))
        return best * size * 0.56 / 72.0 + 0.07

    limits = []
    for i in range(n):
        lim = max([token_in(t["header"][i])] + [token_in(r[i]) for r in t["rows"]])
        limits.append(min(lim, col_width * 0.6))
    widths = [max(w, l) for w, l in zip(widths, limits)]
    spill = sum(widths) - col_width
    if spill > 0:
        slack = [i for i, w in enumerate(widths) if w - limits[i] > 1e-6]
        room = sum(widths[i] - limits[i] for i in slack)
        if room > 0:
            take = min(1.0, spill / room)
            for i in slack:
                widths[i] -= (widths[i] - limits[i]) * take
        scale = col_width / sum(widths)
        widths = [w * scale for w in widths]
    for column, w in zip(tb.columns, widths):
        column.width = _In(w)
    for row in tb.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = _In(w)
    fix_table_layout(tb, col_width)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tb


def add_algo(doc, spec):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    _run(p, spec["title"], "bold", size=8)
    for line in spec["lines"]:
        q = doc.add_paragraph()
        q.paragraph_format.space_before = Pt(0)
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.line_spacing = 1.0
        q.paragraph_format.first_line_indent = Pt(0)
        r = q.add_run(line)
        r.font.size = Pt(7)
        r.font.name = "Consolas"
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_figure(doc, name, caption):
    path = os.path.join(HERE, name)
    if not os.path.exists(path):
        p = doc.add_paragraph()
        _run(p, "[figure %s pending]" % name, size=9)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(3.35))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    m = re.match(r"fig(\d+)", name)
    _run(cap, ("Fig. %s.  " % m.group(1)) if m else "", "bold", size=8)
    add_rich(cap, caption, size=8)
    cap.paragraph_format.space_after = Pt(10)


# ---------------------------------------------------------------- 主流程

def build() -> None:
    doc = Document()
    style_base(doc)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)
    set_one_column(sec)

    # 标题
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    _run(p, C.TITLE, "bold", size=18)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, C.AUTHORS, size=11)
    for line in C.AFFIL.split("\n"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        _run(p, line, size=9, italic=True)

    # 摘要
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    _run(p, "Abstract—", "bold", size=9)
    add_rich(p, C.ABSTRACT, size=9)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    p = doc.add_paragraph()
    _run(p, "Index Terms—", "bold", size=9, italic=True)
    _run(p, C.KEYWORDS, size=9, italic=True)
    p.paragraph_format.left_indent = Inches(0.2)

    # 双栏正文
    body_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_two_columns(body_sec)

    for kind, val in C.BODY:
        if kind == "h1":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(p, val.upper(), "bold", size=10, font="Times New Roman")
        elif kind == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            _run(p, val, "bold", size=10, italic=True)
        elif kind == "p":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_rich(p, val, size=10)
        elif kind == "list":
            for it in val:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                add_rich(p, it, size=10)
        elif kind == "eq":
            add_equation(doc, C.EQ[val], val)
        elif kind == "algo":
            add_algo(doc, val)
        elif kind == "table":
            add_table(doc, BY_NUM[val])
        elif kind == "fig":
            add_figure(doc, *val)

    # 参考文献
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    _run(p, "REFERENCES", "bold", size=10)
    for i, r in enumerate(C.REFS, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        _run(p, "[%d] " % i, size=8)
        add_rich(p, r, size=8)

    doc.save(OUT)
    print("已写出", OUT, "（%.1f KB）" % (os.path.getsize(OUT) / 1024))


if __name__ == "__main__":
    build()
